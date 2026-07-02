import glob
import json
import math
import os
import re
from collections import Counter

ROOT = os.path.dirname(os.path.abspath(__file__))
DOCS_DIR = os.path.join(ROOT, "proper_docs")
PDF_DIR = os.path.join(ROOT, "proper_docs_pdf")
OUT = os.path.join(ROOT, "data", "template_store.json")

PAGES_OVERRIDE = {"GCL Secure DOGH": 2, "GCL Secure_LFQ": 2}


def page_count(name):
    pdf = os.path.join(PDF_DIR, name + ".pdf")
    if os.path.exists(pdf):
        from pdf2image import pdfinfo_from_path
        return int(pdfinfo_from_path(pdf)["Pages"])
    if name in PAGES_OVERRIDE:
        return PAGES_OVERRIDE[name]
    raise SystemExit(f"no page count for template {name!r}: cache a PDF in {PDF_DIR} or add to PAGES_OVERRIDE")


def tokens(s):
    return [t for t in re.sub(r"[^a-z0-9]", " ", str(s).lower()).split() if len(t) >= 2]


def docx_tokens(path):
    import docx
    d = docx.Document(path)
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        for row in t.rows:
            seen = None
            for c in row.cells:
                txt = c.text.strip()
                if txt and txt != seen:      # drop merge-cell duplicates repeated across a span
                    parts.append(txt)
                seen = txt
    return tokens(" ".join(parts))


def tf_vector(toks, vocab, idf):
    counts = Counter(t for t in toks if t in vocab)
    vec = [0.0] * len(vocab)
    for tok, n in counts.items():
        vec[vocab[tok]] = (1.0 + math.log(n)) * idf[vocab[tok]]
    norm = math.sqrt(sum(v * v for v in vec))
    if norm:
        vec = [v / norm for v in vec]
    return vec


def build():
    paths = sorted(glob.glob(os.path.join(DOCS_DIR, "*.docx")))
    if not paths:
        raise SystemExit(f"no .docx templates in {DOCS_DIR}")
    doc_toks = {os.path.splitext(os.path.basename(p))[0]: docx_tokens(p) for p in paths}

    df = Counter()
    for toks in doc_toks.values():
        df.update(set(toks))
    vocab = {tok: i for i, tok in enumerate(sorted(df))}
    n = len(doc_toks)
    idf = [0.0] * len(vocab)
    for tok, i in vocab.items():
        idf[i] = math.log((1 + n) / (1 + df[tok])) + 1.0

    docs = [
        {"name": name, "vec": tf_vector(toks, vocab, idf), "n_pages": page_count(name)}
        for name, toks in doc_toks.items()
    ]
    store = {"vocab": vocab, "idf": idf, "docs": docs}
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    with open(OUT, "w", encoding="utf-8") as f:
        json.dump(store, f)
    print(f"built {OUT}: {n} templates, vocab={len(vocab)}")
    for d in docs:
        print(f"  {d['name']}  n_pages={d['n_pages']}")
    return store


if __name__ == "__main__":
    build()
