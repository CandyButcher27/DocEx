import csv
import json

from src.exporters.exporter import export_document
from src.exporters.serializer import document_to_records
from src.models.document import Document, ExtractionRoute
from src.models.field import Field
from src.models.ocr_result import OCRResult
from src.models.page import Page
from src.models.validation_result import ValidationResult
from src.template.enums import ExportFormat, FieldType, FieldWidget, OCREngine, ValidatorType
from src.template.export_config import ExportConfig
from src.template.field_config import FieldConfig
from src.template.ocr_config import OCRConfig
from src.template.roi_config import ROIConfig
from src.template.validation_config import ValidationConfig


def _field(field_id: str, value: str) -> Field:
    config = FieldConfig(
        id=field_id,
        label=field_id.replace("_", " ").title(),
        datatype=FieldType.TEXT,
        widget=FieldWidget.TEXTBOX,
        roi=ROIConfig(x=0.1, y=0.1, width=0.2, height=0.1),
        ocr=OCRConfig(engine=OCREngine.PADDLE),
        validation=ValidationConfig(validator=ValidatorType.TEXT),
        export=ExportConfig(key=field_id, format=ExportFormat.JSON),
    )
    return Field(
        config=config,
        ocr_result=OCRResult(raw_text=value, confidence=0.9, engine_name="stub", processing_time_ms=1.0),
        validation_result=ValidationResult(normalized_text=value, is_valid=True, validation_score=1.0),
    )


def _document() -> Document:
    page = Page(number=1, fields=[_field("name", "John"), _field("city", "Pune")])
    return Document(source_path="x.pdf", template_id="t", route=ExtractionRoute.TEMPLATE_MATCH, pages=[page])


def test_records_shape():
    records = document_to_records(_document())
    assert len(records) == 2
    assert records[0]["key"] == "name"
    assert records[0]["value"] == "John"
    assert records[0]["confidence"] == 0.95


def test_export_json(tmp_path):
    out = export_document(_document(), ExportFormat.JSON, tmp_path / "out")
    assert out.suffix == ".json"
    data = json.loads(out.read_text(encoding="utf-8"))
    assert len(data["fields"]) == 2


def test_export_csv(tmp_path):
    out = export_document(_document(), ExportFormat.CSV, tmp_path / "out")
    assert out.suffix == ".csv"
    rows = list(csv.DictReader(out.open(encoding="utf-8")))
    assert len(rows) == 2
    assert rows[0]["value"] == "John"


def test_export_excel(tmp_path):
    from openpyxl import load_workbook

    out = export_document(_document(), ExportFormat.EXCEL, tmp_path / "out")
    assert out.suffix == ".xlsx"
    ws = load_workbook(out).active
    assert ws.max_row == 3
